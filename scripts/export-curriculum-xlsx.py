#!/usr/bin/env python3
"""Export curriculum data from PostgreSQL to XLSX.

Usage:
    python scripts/export-curriculum-xlsx.py
    python scripts/export-curriculum-xlsx.py --database-url postgresql://... --output ./data/curriculum-export.xlsx

Requires: psycopg[binary], openpyxl
Env:      DB_URL (PostgreSQL connection string, used as default for --database-url)
"""

from __future__ import annotations

import argparse
import os
import sys
from datetime import datetime, timezone
from pathlib import Path

import psycopg
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

GENERATOR_NAME = "curriculum-export"
GENERATOR_VERSION = "1.0.0"

# ── Query definitions ────────────────────────────────────────────────────────

QUERIES = {
    "Curriculum": {
        "columns": [
            "grade_key", "grade", "subject", "category", "seq_number",
            "lesson_topic", "skill_statement", "standards", "song_count",
            "linked_songs", "linked_resources", "tags", "circle_time_slot",
        ],
        "sql": (
            "SELECT grade_key, grade, subject, category, seq_number,"
            "       lesson_topic, skill_statement, standards, song_count,"
            "       linked_songs, linked_resources, tags, circle_time_slot"
            "  FROM curriculum_topics"
            "  ORDER BY grade_key, seq_number"
        ),
    },
    "Songs": {
        "columns": [
            "id", "title", "artist", "type", "age_range",
            "educational_domain", "lyrics", "instructions", "actions",
            "materials_needed", "tags", "verified", "curriculum_links",
        ],
        "sql": (
            "SELECT id, title, artist, type, age_range,"
            "       educational_domain,"
            "       LEFT(lyrics, 200) AS lyrics,"
            "       instructions, actions, materials_needed,"
            "       tags, verified, curriculum_links"
            "  FROM songs"
            "  ORDER BY id"
        ),
    },
    "Activities": {
        "columns": [
            "id", "name", "type", "instructions",
            "materials_needed", "age_range", "duration_minutes",
        ],
        "sql": (
            "SELECT id, name, type, instructions,"
            "       materials_needed, age_range, duration_minutes"
            "  FROM activities"
            "  ORDER BY id"
        ),
    },
    "Song Actions": {
        "columns": [
            "id", "song_title", "alternate_title", "tradition_performer",
            "action_wording", "normalized_action", "action_sequence",
            "action_classification", "core_or_optional",
            "age_range_stated", "educator_org", "source_title",
            "research_status",
        ],
        "sql": (
            "SELECT id, song_title, alternate_title, tradition_performer,"
            "       action_wording, action AS normalized_action,"
            "       action_sequence, action_classification,"
            "       core_or_optional, age_range_stated,"
            "       educator_org, source_title, research_status"
            "  FROM song_actions"
            "  ORDER BY id"
        ),
    },
}

# ── Helpers ──────────────────────────────────────────────────────────────────

MAX_COL_WIDTH = 60


def auto_size_columns(ws) -> None:
    """Set each column width to max content width, capped at MAX_COL_WIDTH."""
    for col_idx, col_cells in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in col_cells:
            val = str(cell.value) if cell.value is not None else ""
            max_len = max(max_len, len(val))
        # +2 for padding
        width = min(max_len + 2, MAX_COL_WIDTH)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def write_sheet(wb: Workbook, name: str, columns: list[str], rows: list[tuple]) -> int:
    """Write a data sheet. Returns the number of data rows written."""
    ws = wb.create_sheet(title=name)
    # Header row
    ws.append(columns)
    for cell in ws[1]:
        cell.alignment = Alignment(wrap_text=True)

    for row in rows:
        ws.append(list(row))

    auto_size_columns(ws)
    return len(rows)


def write_metadata_sheet(wb: Workbook, row_counts: dict[str, int], timestamp: str) -> None:
    """Write the Metadata sheet with generation info and row counts."""
    ws = wb.create_sheet(title="Metadata")
    ws.append(["Key", "Value"])
    ws.append(["generator_name", GENERATOR_NAME])
    ws.append(["generator_version", GENERATOR_VERSION])
    ws.append(["export_timestamp", timestamp])
    ws.append([])  # blank separator row
    for sheet_name, count in row_counts.items():
        ws.append([f"row_count_{sheet_name.lower().replace(' ', '_')}", count])
    auto_size_columns(ws)


# ── Main ─────────────────────────────────────────────────────────────────────

def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Export curriculum data to XLSX")
    p.add_argument(
        "--database-url",
        default=os.environ.get("DB_URL"),
        help="PostgreSQL connection URL (default: $DB_URL)",
    )
    p.add_argument(
        "--output",
        default="./data/curriculum-export.xlsx",
        help="Output XLSX path (default: ./data/curriculum-export.xlsx)",
    )
    args = p.parse_args(argv)
    if not args.database_url:
        p.error("No database URL. Set DB_URL or pass --database-url.")
    return args


def main(argv: list[str] | None = None) -> None:
    args = parse_args(argv)

    timestamp = datetime.now(timezone.utc).isoformat()
    row_counts: dict[str, int] = {}

    out_path = Path(args.output)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    with psycopg.connect(
        args.database_url, options="-c default_transaction_read_only=on"
    ) as conn:
        with conn.cursor() as cur:
            wb = Workbook()
            # Remove the default sheet created by openpyxl
            wb.remove(wb.active)

            for sheet_name, spec in QUERIES.items():
                cur.execute(spec["sql"])
                columns = spec["columns"]
                rows = cur.fetchall()
                count = write_sheet(wb, sheet_name, columns, rows)
                row_counts[sheet_name] = count
                print(f"  {sheet_name}: {count} rows")

            write_metadata_sheet(wb, row_counts, timestamp)

    wb.save(args.output)
    print(f"\nExported to {out_path}")
    print(f"Total sheets: {len(QUERIES) + 1}")


if __name__ == "__main__":
    main()